"""Downloading a chat result as a file: Excel and PDF.

The rows already left the server as JSON (the snapshot rides on the answer),
so this endpoint changes the FORMAT of that same egress, not what egresses.
Scoping is identical to /runs/{id}: owner-only, 404 otherwise.
"""
import io

import pytest
from openpyxl import load_workbook

from app.models.models import AgentMessage, AgentRun, AgentStep, DataSource

SNAPSHOT = {"columns": ["city", "n"], "rows": [["Cairo", 3], ["Giza", 1]],
            "total": 2, "truncated": False}
SQL = "SELECT city, count(*) AS n FROM t GROUP BY city"


@pytest.fixture
def scripted(monkeypatch):
    async def fake_run(db, *, question, user, client, source=None,
                       datasets=None, conversation_id=None, history=None, **kw):
        run = AgentRun(org_id=user.org_id, conversation_id=conversation_id,
                       question=question, status="ok", intent="lookup",
                       answer="Cairo leads with 3.")
        db.add(run)
        await db.flush()
        db.add(AgentStep(agent_run_id=run.id, node="s1", status="ok",
                         sql=SQL, rows_returned=2, result_rows=SNAPSHOT))
        await db.flush()
        return run
    monkeypatch.setattr("app.routers.agent.run_agent", fake_run)
    return fake_run


@pytest.fixture
async def source(db_session, two_orgs):
    src = DataSource(name="wh-a", type="postgresql", org_id=two_orgs["a"]["org"].id)
    db_session.add(src)
    await db_session.commit()
    return src


@pytest.fixture
async def answered_run(client, auth_headers, source, scripted):
    """A conversation with one answered question; yields the run id."""
    r = await client.post("/api/v1/agent/conversations",
                          json={"data_source_id": source.id},
                          headers=auth_headers["a"])
    cid = r.json()["id"]
    got = await client.post(f"/api/v1/agent/conversations/{cid}/ask",
                            json={"question": "orders per city"},
                            headers=auth_headers["a"])
    return {"cid": cid, "run_id": got.json()["run_id"]}


class TestExcel:
    async def test_the_workbook_holds_the_snapshot(self, client, auth_headers,
                                                   answered_run):
        r = await client.get(
            f"/api/v1/agent/runs/{answered_run['run_id']}/export?format=xlsx",
            headers=auth_headers["a"])
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml")
        assert f"ask-ai-result-{answered_run['run_id']}.xlsx" in \
            r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb[wb.sheetnames[0]]
        assert [c.value for c in ws[1]] == ["city", "n"]
        assert [c.value for c in ws[2]] == ["Cairo", 3]
        assert [c.value for c in ws[3]] == ["Giza", 1]


class TestPdf:
    async def test_a_pdf_comes_back(self, client, auth_headers, answered_run):
        r = await client.get(
            f"/api/v1/agent/runs/{answered_run['run_id']}/export?format=pdf",
            headers=auth_headers["a"])
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"
        assert r.content[:5] == b"%PDF-"
        assert f"ask-ai-result-{answered_run['run_id']}.pdf" in \
            r.headers["content-disposition"]


class TestGuards:
    async def test_an_unknown_format_is_400(self, client, auth_headers, answered_run):
        r = await client.get(
            f"/api/v1/agent/runs/{answered_run['run_id']}/export?format=docx",
            headers=auth_headers["a"])
        assert r.status_code == 400

    async def test_a_run_with_no_rows_is_400(self, client, auth_headers,
                                             db_session, two_orgs, source):
        # A clarification (or a failure) has nothing tabular to export.
        from app.models.models import Conversation
        conv = Conversation(org_id=two_orgs["a"]["org"].id,
                            user_id=two_orgs["a"]["user"].id,
                            data_source_id=source.id, title="t")
        db_session.add(conv)
        await db_session.flush()
        run = AgentRun(org_id=two_orgs["a"]["org"].id, conversation_id=conv.id,
                       question="q", status="needs_clarification")
        db_session.add(run)
        await db_session.commit()
        r = await client.get(f"/api/v1/agent/runs/{run.id}/export?format=xlsx",
                             headers=auth_headers["a"])
        assert r.status_code == 400

    async def test_owner_only(self, client, auth_headers, answered_run,
                              db_session, two_orgs):
        from app.core.security import create_access_token, hash_password
        from app.models.models import User
        org_a = two_orgs["a"]["org"]
        second = User(org_id=org_a.id, role_id=two_orgs["a"]["role"].id,
                      email="second@example.com", password_hash=hash_password("pw"))
        db_session.add(second)
        await db_session.commit()
        headers = {"Authorization":
                   f"Bearer {create_access_token(second.id, second.org_id)}"}
        for h in (headers, auth_headers["b"]):
            r = await client.get(
                f"/api/v1/agent/runs/{answered_run['run_id']}/export?format=xlsx",
                headers=h)
            assert r.status_code == 404


class TestPdfBuilder:
    def test_arabic_cells_are_shaped_and_get_a_capable_font(self):
        """The builder is pure; pin the Arabic path directly: shaped text and
        the DejaVu font must both appear in the document, or Arabic cells
        draw as blank space (Helvetica has no Arabic glyphs)."""
        from app.services.agent.export import build_result_pdf
        arabic_snap = {"columns": ["المدينة", "n"],
                       "rows": [["القاهرة", 3]], "total": 1, "truncated": False}
        pdf = build_result_pdf("كم طلب لكل مدينة؟", "القاهرة 3.",
                               [SQL], [("s1", arabic_snap)])
        assert pdf[:5] == b"%PDF-"
        assert b"DejaVuSans" in pdf

    def test_a_truncated_snapshot_says_so(self):
        # The wording is pinned on its own helper (PDF streams are
        # compressed, so grepping the binary would test the compressor);
        # the builder must also accept such a snapshot without raising.
        from app.services.agent.export import build_result_pdf, trunc_note
        assert trunc_note(1, 500) == "Showing the first 1 of 500 rows."
        snap = {"columns": ["n"], "rows": [[1]], "total": 500, "truncated": True}
        assert build_result_pdf("q", None, [], [("s1", snap)])[:5] == b"%PDF-"

    def test_sql_with_angle_brackets_survives(self):
        # Paragraph parses XML-ish markup; unescaped "<>" would raise deep in
        # reportlab and kill the export.
        from app.services.agent.export import build_result_pdf
        pdf = build_result_pdf("q", "a", ["SELECT 1 WHERE a <> b AND c < 2"],
                               [("s1", SNAPSHOT)])
        assert pdf[:5] == b"%PDF-"
